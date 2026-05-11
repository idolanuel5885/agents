"""Mock-based tests for the radius-based Govmap deals integration.

Covers:
* :func:`govmap_client.autocomplete_address` — parse, filter, fail
* :func:`govmap_client.find_comparable_deals` — full pipeline including
  apartment filter, polygon-level address enrichment, sort+limit,
  partial-success on per-polygon failure, empty radius
* Web endpoints ``/shuma/autocomplete`` and ``/shuma/comparables``
* End-to-end report rendering: outlier marker + balcony-column hiding
* A skipped live-API smoke test
"""
from __future__ import annotations

import io
import os
import sys
from unittest.mock import patch, MagicMock

import pytest


ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from real_estate import govmap_client  # noqa: E402
from real_estate.govmap_client import (  # noqa: E402
    GovmapFetchError,
    autocomplete_address,
    find_comparable_deals,
    _format_floor,
    _format_deal_date,
)
from real_estate.models import ComparisonProperty  # noqa: E402


# ── Helpers ──────────────────────────────────────────────────────────────────


def _mock_response(status: int = 200, body=None, raise_on_json: bool = False):
    resp = MagicMock()
    resp.status_code = status
    resp.ok = 200 <= status < 400
    resp.content = b"x"
    resp.text = "x"
    if raise_on_json:
        resp.json.side_effect = ValueError("not json")
    else:
        resp.json.return_value = body if body is not None else {}
    return resp


@pytest.fixture(autouse=True)
def _no_sleep():
    """Skip the 0.2 s polite delay between street-deals calls in tests."""
    with patch("real_estate.govmap_client.time.sleep"):
        yield


# ── 1. autocomplete_address ──────────────────────────────────────────────────


def test_autocomplete_returns_parsed_results():
    body = {
        "resultsCount": 2,
        "results": [
            {
                "id": "address|ADDRESS|64834989|רוטשילד 1|תל אביב",
                "text": "רוטשילד 1, תל אביב",
                "type": "address",
                "shape": "POINT(3871406.911 3772053.769)",
            },
            {
                "id": "street|STREET_MID_POINT|36595|רוטשילד|תל אביב",
                "text": "רוטשילד, תל אביב",
                "type": "street",
                "shape": "POINT(3871500.0 3772100.0)",
            },
        ],
    }
    with patch.object(govmap_client._session, "post", return_value=_mock_response(body=body)):
        out = autocomplete_address("רוטשילד")

    assert len(out) == 2
    assert out[0]["display_name"] == "רוטשילד 1, תל אביב"
    assert out[0]["type"] == "address"
    assert out[0]["addr_id"] == "64834989"
    assert out[0]["itm_x"] == 3871406.911
    assert out[0]["itm_y"] == 3772053.769
    assert out[1]["type"] == "street"
    assert out[1]["addr_id"] is None  # street rows don't carry addr_id


def test_autocomplete_short_query_short_circuits():
    with patch.object(govmap_client._session, "post") as p:
        out = autocomplete_address("ר")
    assert out == []
    p.assert_not_called()


def test_autocomplete_500_raises_hebrew_error():
    with patch.object(govmap_client._session, "post", return_value=_mock_response(status=500)):
        with pytest.raises(GovmapFetchError) as exc:
            autocomplete_address("רוטשילד")
    assert exc.value.code == "AUTOCOMPLETE_UNAVAILABLE"
    assert "השלמת" in exc.value.message_he


def test_autocomplete_connection_error_raises_hebrew():
    with patch.object(govmap_client._session, "post", side_effect=ConnectionError("dns")):
        with pytest.raises(GovmapFetchError) as exc:
            autocomplete_address("רוטשילד")
    assert exc.value.code == "AUTOCOMPLETE_UNAVAILABLE"


# ── 2. find_comparable_deals — full pipeline ─────────────────────────────────


def _polygon(polygon_id, street, house, dealscount=5):
    return {
        "polygon_id": polygon_id,
        "streetNameHeb": street,
        "houseNum": house,
        "dealscount": dealscount,
    }


def _deal(
    *,
    polygon_id="7422-116",
    date_="2025-02-15T00:00:00.000Z",
    price=4_500_000,
    area=110,
    rooms=4,
    floor=5,
    property_type="דירה",
    neighborhood="לב העיר",
    settlement="תל אביב -יפו",
):
    return {
        "polygonId": polygon_id,
        "streetNameHeb": None,
        "houseNum": None,
        "assetRoomNum": rooms,
        "floorNo": floor,
        "assetArea": area,
        "dealAmount": price,
        "dealDate": date_,
        "propertyTypeDescription": property_type,
        "neighborhood": neighborhood,
        "settlementNameHeb": settlement,
    }


def _build_session_mock(radius_body, street_deals_by_polygon):
    """Build a MagicMock session whose .get() routes by URL substring."""
    def _get(url, **kwargs):
        if "/real-estate/deals/" in url and "street-deals" not in url:
            return _mock_response(body=radius_body)
        for pid, body in street_deals_by_polygon.items():
            if f"/street-deals/{pid}" in url:
                return _mock_response(body=body)
        return _mock_response(status=404)
    return _get


def test_find_comparable_deals_full_pipeline_with_enrichment():
    """Polygon-level address must override the deal's null streetNameHeb."""
    polygons = [
        _polygon("7422-116", "רוטשילד", "12", dealscount=20),
    ]
    deals = {
        "7422-116": {"data": [_deal(price=4_500_000, area=110, rooms=4, floor=5)]},
    }

    with patch.object(govmap_client._session, "get", side_effect=_build_session_mock(polygons, deals)):
        out = find_comparable_deals(3871406.911, 3772053.769, radius_m=200, max_deals=10)

    assert len(out) == 1
    cp = out[0]
    # Address comes from the polygon, not the deal (which had null)
    assert cp.address == "רוטשילד 12"
    assert cp.rooms == "4"
    assert cp.floor == "5"
    assert cp.built_area == 110.0
    assert cp.price == 4_500_000.0
    assert cp.balcony_area is None
    assert cp.notes == "15.2.2025"
    assert cp.is_outlier is False


def test_find_comparable_deals_filters_non_apartments():
    """Stores, plots, buildings — none of them pass the apartment filter."""
    polygons = [_polygon("7422-116", "רוטשילד", "12", dealscount=20)]
    deals = {
        "7422-116": {
            "data": [
                _deal(property_type="חנות", price=1, area=1, rooms=0),
                _deal(property_type="בנין", price=2, area=2, rooms=0),
                _deal(property_type="מגרש", price=3, area=3, rooms=0),
                _deal(property_type="דירה", price=4_500_000, area=110, rooms=4),
            ],
        },
    }
    with patch.object(govmap_client._session, "get", side_effect=_build_session_mock(polygons, deals)):
        out = find_comparable_deals(0.0, 0.0)
    assert len(out) == 1
    assert out[0].built_area == 110.0


def test_find_comparable_deals_sorts_desc_and_limits():
    polygons = [_polygon("p1", "רחוב א", "1", dealscount=99)]
    deals = {
        "p1": {
            "data": [
                _deal(date_="2023-01-01T00:00:00Z", price=1, area=10),
                _deal(date_="2025-06-30T00:00:00Z", price=2, area=10),
                _deal(date_="2024-09-15T00:00:00Z", price=3, area=10),
                _deal(date_="2025-12-01T00:00:00Z", price=4, area=10),
            ],
        },
    }
    with patch.object(govmap_client._session, "get", side_effect=_build_session_mock(polygons, deals)):
        out = find_comparable_deals(0.0, 0.0, radius_m=200, max_deals=2)
    assert len(out) == 2
    assert out[0].notes == "1.12.2025"  # newest first
    assert out[1].notes == "30.6.2025"


def test_find_comparable_deals_partial_success_on_polygon_failure():
    """One flaky polygon shouldn't blank the whole table."""
    polygons = [
        _polygon("good", "רחוב א", "1", dealscount=99),
        _polygon("bad",  "רחוב ב", "2", dealscount=80),
    ]
    # 'bad' polygon raises a connection error mid-call — should be skipped silently.
    def _get(url, **kwargs):
        if "street-deals/bad" in url:
            raise ConnectionError("flaky")
        if "street-deals/good" in url:
            return _mock_response(body={"data": [_deal()]})
        if "/real-estate/deals/" in url:
            return _mock_response(body=polygons)
        return _mock_response(status=404)

    with patch.object(govmap_client._session, "get", side_effect=_get):
        out = find_comparable_deals(0.0, 0.0)

    # 'good' polygon's deals come through; 'bad' is silently dropped.
    assert len(out) == 1
    assert out[0].address.startswith("רחוב א")


def test_find_comparable_deals_no_polygons_returns_empty():
    with patch.object(govmap_client._session, "get", return_value=_mock_response(body=[])):
        out = find_comparable_deals(0.0, 0.0)
    assert out == []


def test_find_comparable_deals_radius_500_raises():
    with patch.object(govmap_client._session, "get", return_value=_mock_response(status=502)):
        with pytest.raises(GovmapFetchError) as exc:
            find_comparable_deals(0.0, 0.0)
    assert exc.value.code == "POLYGONS_FETCH_FAILED"


def test_find_comparable_deals_caps_top_polygons_by_dealscount():
    """Only the top _MAX_POLYGONS (10) are queried; sorted by dealscount."""
    polygons = [_polygon(f"p{i}", f"רחוב {i}", "1", dealscount=i) for i in range(15)]
    deals = {f"p{i}": {"data": [_deal(date_=f"2024-{i+1:02d}-01T00:00:00Z")]} for i in range(15)}
    queried_polygons: list[str] = []

    def _get(url, **kwargs):
        if "/real-estate/deals/" in url and "street-deals" not in url:
            return _mock_response(body=polygons)
        # Match the polygon id as the URL's trailing path segment so "p1"
        # doesn't get confused with "p10", "p11", etc.
        pid = url.rsplit("/", 1)[-1]
        if pid in deals:
            queried_polygons.append(pid)
            return _mock_response(body=deals[pid])
        return _mock_response(status=404)

    with patch.object(govmap_client._session, "get", side_effect=_get):
        find_comparable_deals(0.0, 0.0)

    # We queried exactly _MAX_POLYGONS (10) — the highest-dealscount ones.
    assert len(queried_polygons) == 10
    # p14 (dealscount=14) made the cut; p0 (dealscount=0) did not.
    assert "p14" in queried_polygons
    assert "p0" not in queried_polygons


def test_find_comparable_deals_falls_back_to_neighborhood_when_no_street():
    """If the polygon also has null streetNameHeb, render neighborhood + city."""
    polygons = [{
        "polygon_id": "p1",
        "streetNameHeb": None,
        "houseNum": None,
        "dealscount": 5,
    }]
    deals = {"p1": {"data": [_deal()]}}

    with patch.object(govmap_client._session, "get", side_effect=_build_session_mock(polygons, deals)):
        out = find_comparable_deals(0.0, 0.0)

    assert len(out) == 1
    assert "לב העיר" in out[0].address
    assert "תל אביב" in out[0].address


# ── 3. Field-level helpers ───────────────────────────────────────────────────


def test_format_floor_translates_hebrew_words():
    assert _format_floor("אחת עשרה") == "11"
    assert _format_floor("שתיים עשרה") == "12"
    assert _format_floor("עשרים") == "20"
    assert _format_floor("קומת קרקע") == "0"
    assert _format_floor("ראשונה") == "1"
    # Unknown Hebrew word → "" (appraiser fills, never fabricated).
    assert _format_floor("לא מוכר") == ""
    # Numeric pass-through.
    assert _format_floor(7) == "7"
    assert _format_floor("3") == "3"


def test_format_deal_date_israeli_convention():
    assert _format_deal_date("2024-06-19T00:00:00.000Z") == "19.6.2024"
    assert _format_deal_date("2015-02-08T00:00:00Z") == "8.2.2015"
    assert _format_deal_date("") == ""
    assert _format_deal_date(None) == ""
    assert _format_deal_date("garbage") == ""


# ── 4. Web endpoints ─────────────────────────────────────────────────────────


@pytest.fixture
def client():
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from real_estate.web import router
    app = FastAPI()
    app.include_router(router)
    return TestClient(app)


def test_autocomplete_endpoint_success(client):
    with patch("real_estate.govmap_client.autocomplete_address") as p:
        p.return_value = [{
            "display_name": "x", "type": "address",
            "itm_x": 1.0, "itm_y": 2.0, "addr_id": "9",
        }]
        r = client.post("/shuma/autocomplete", json={"query": "רוטשילד"})
    assert r.status_code == 200
    body = r.json()
    assert body["success"] is True
    assert body["results"][0]["itm_x"] == 1.0


def test_autocomplete_endpoint_failure_returns_hebrew(client):
    with patch("real_estate.govmap_client.autocomplete_address") as p:
        p.side_effect = GovmapFetchError("AUTOCOMPLETE_UNAVAILABLE", "שירות לא זמין")
        r = client.post("/shuma/autocomplete", json={"query": "רוטשילד"})
    body = r.json()
    assert r.status_code == 200
    assert body["success"] is False
    assert body["error_code"] == "AUTOCOMPLETE_UNAVAILABLE"


def test_comparables_endpoint_success(client):
    deals = [ComparisonProperty(
        address="רוטשילד 12", floor="5", rooms="4",
        built_area=110.0, balcony_area=None, price=4_500_000,
        notes="15.2.2025", is_outlier=False,
    )]
    with patch("real_estate.govmap_client.find_comparable_deals", return_value=deals):
        r = client.post(
            "/shuma/comparables",
            json={"itm_x": 3871406.911, "itm_y": 3772053.769, "radius_m": 200, "max_deals": 10},
        )
    assert r.status_code == 200
    body = r.json()
    assert body["success"] is True
    assert len(body["deals"]) == 1
    assert body["deals"][0]["address"] == "רוטשילד 12"
    assert body["fetched_at"]
    assert body["source"] == "govmap.gov.il"


def test_comparables_endpoint_failure(client):
    with patch("real_estate.govmap_client.find_comparable_deals") as p:
        p.side_effect = GovmapFetchError("POLYGONS_FETCH_FAILED", "שירות לא זמין")
        r = client.post(
            "/shuma/comparables",
            json={"itm_x": 1.0, "itm_y": 2.0},
        )
    body = r.json()
    assert r.status_code == 200
    assert body["success"] is False
    assert body["error_code"] == "POLYGONS_FETCH_FAILED"


def test_comparables_endpoint_empty_returns_no_deals_found(client):
    with patch("real_estate.govmap_client.find_comparable_deals", return_value=[]):
        r = client.post(
            "/shuma/comparables",
            json={"itm_x": 1.0, "itm_y": 2.0, "radius_m": 200},
        )
    body = r.json()
    assert body["success"] is True
    assert body["error_code"] == "NO_DEALS_FOUND"
    assert body["deals"] == []


# ── 5. End-to-end report (outlier marker + balcony hiding) ───────────────────


def test_report_renders_outlier_marker_and_hides_balcony_columns():
    from real_estate.demo_data import sample_market
    from real_estate.report_generator import generate_report_bytes
    from docx import Document

    data = sample_market()
    data.comparables_fetched_at = "2025-11-12"
    data.comparison_properties = [
        ComparisonProperty(
            address="רוטשילד 5, תל אביב", floor="3", rooms="3",
            built_area=80.0, balcony_area=None, price=3_000_000,
            notes="11.2.2025", is_outlier=True,
        ),
        ComparisonProperty(
            address="רוטשילד 7, תל אביב", floor="2", rooms="3",
            built_area=85.0, balcony_area=None, price=3_200_000,
            notes="4.11.2024", is_outlier=False,
        ),
        ComparisonProperty(
            address="רוטשילד 9, תל אביב", floor="4", rooms="3",
            built_area=82.0, balcony_area=12.0, price=3_350_000,
            notes="3.10.2024", is_outlier=False,
        ),
    ]

    docx_bytes = generate_report_bytes(data)
    doc = Document(io.BytesIO(docx_bytes))
    text = "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(
        cell.text for table in doc.tables for row in table.rows for cell in row.cells
    )

    assert "(*) 1" in text
    assert "חריגות ביחס למקובל" in text
    assert "2025-11-12" in text
    table_headers = [
        cell.text for table in doc.tables for cell in table.rows[0].cells
    ]
    assert "מרפסת מ\"ר" not in table_headers


# ── 6. Live smoke (skipped by default) ───────────────────────────────────────


@pytest.mark.skip(reason="Live API test — run manually after deploy")
def test_live_endpoint_smoke():  # pragma: no cover
    """Manual probe: autocomplete a real address and pull comparables.

    Expectation when run from a place that can actually reach Govmap:
    at least 5 apartment deals from multiple streets, with dealDate
    in the last 2 years.
    """
    results = autocomplete_address("רוטשילד 1 תל אביב")
    assert results
    addr = next(r for r in results if r["type"] == "address")
    deals = find_comparable_deals(addr["itm_x"], addr["itm_y"], radius_m=200, max_deals=10)
    assert len(deals) >= 5
    streets = {d.address.split()[0] for d in deals if d.address}
    assert len(streets) >= 2, f"only one street in deals: {streets}"
