"""RTL-aware python-docx helpers for Hebrew documents.

הקובץ הזה מרכז את כל ההתאמות של python-docx לעברית RTL — גודל עמוד A4,
שוליים, גופן David, bidiVisual לטבלאות, header/footer של עטר לרון, ועוד.
כל ייצור הדוחות חייב לעבור דרך העזרים האלה — לא לכתוב XML ישירות
בקוד הסקשנים.
"""
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# גופן ברירת מחדל — David, גודלים בנקודות (Pt), כפי שמופיעים בקובץ ה-XML *2
HEBREW_FONT = "David"
FONT_BODY = 12
FONT_HEADING2 = 13
FONT_HEADING1 = 14
FONT_TITLE = 16

# הנכסים של תיקיית האסטים — לוגו עטר לרון ו-footer.
_ASSETS_DIR = Path(__file__).parent / "assets"
_HEADER_LOGO = _ASSETS_DIR / "header_logo.png"
_FOOTER_STRIP = _ASSETS_DIR / "footer_strip.png"


# ── Document setup ────────────────────────────────────────────────────────────

def make_rtl_doc() -> Document:
    """Create an A4 document configured for Hebrew RTL output."""
    doc = Document()
    section = doc.sections[0]
    _configure_section_a4(section)
    _set_doc_rtl_defaults(doc)
    _set_settings_hebrew(doc)
    _attach_header_footer(doc)
    return doc


def _configure_section_a4(section):
    """A4 page size, margins, bidi/rtlGutter/titlePg on sectPr."""
    sectPr = section._sectPr

    # מסירים pgSz/pgMar קיימים אם יש (python-docx מציב כברירת מחדל Letter)
    for tag in ("w:pgSz", "w:pgMar", "w:bidi", "w:rtlGutter", "w:titlePg"):
        for el in sectPr.findall(qn(tag)):
            sectPr.remove(el)

    pgSz = OxmlElement("w:pgSz")
    pgSz.set(qn("w:w"), "11906")   # A4 width in DXA (1/20 pt)
    pgSz.set(qn("w:h"), "16838")   # A4 height in DXA
    sectPr.append(pgSz)

    pgMar = OxmlElement("w:pgMar")
    pgMar.set(qn("w:top"), "1440")
    pgMar.set(qn("w:right"), "1800")
    pgMar.set(qn("w:bottom"), "1440")
    pgMar.set(qn("w:left"), "1800")
    pgMar.set(qn("w:header"), "708")
    pgMar.set(qn("w:footer"), "1020")
    pgMar.set(qn("w:gutter"), "0")
    sectPr.append(pgMar)

    bidi = OxmlElement("w:bidi")
    sectPr.append(bidi)

    rtlGutter = OxmlElement("w:rtlGutter")
    sectPr.append(rtlGutter)

    titlePg = OxmlElement("w:titlePg")
    sectPr.append(titlePg)


def _set_doc_rtl_defaults(doc: Document):
    """Force the Normal style to RTL + David at body size."""
    normal = doc.styles["Normal"]
    pPr = normal.element.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    pPr.append(bidi)
    jc = OxmlElement("w:jc")
    jc.set(qn("w:val"), "right")
    pPr.append(jc)

    rPr = normal.element.get_or_add_rPr()
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), HEBREW_FONT)
    rFonts.set(qn("w:hAnsi"), HEBREW_FONT)
    rFonts.set(qn("w:cs"), HEBREW_FONT)
    rPr.append(rFonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(FONT_BODY * 2))
    rPr.append(sz)
    szCs = OxmlElement("w:szCs")
    szCs.set(qn("w:val"), str(FONT_BODY * 2))
    rPr.append(szCs)


def _set_settings_hebrew(doc: Document):
    """Mark the document language as Hebrew (he-IL bidi) in settings.xml."""
    settings = doc.settings.element
    # מוחקים themeFontLang קיימים כדי למנוע כפילות
    for el in settings.findall(qn("w:themeFontLang")):
        settings.remove(el)
    lang = OxmlElement("w:themeFontLang")
    lang.set(qn("w:val"), "en-US")
    lang.set(qn("w:bidi"), "he-IL")
    settings.append(lang)


# ── Header / Footer ───────────────────────────────────────────────────────────

def _attach_header_footer(doc: Document):
    """Insert the Atar-Laron logo header and contact-strip footer.

    משתמש ב-titlePg כדי להחזיק תשתית ל-first-page header נפרד בעתיד; כרגע
    כל העמודים מציגים את אותו ה-header וה-footer. אם תמונת ה-asset חסרה
    (למשל בסביבת בדיקות), הפסקה נוצרת ריקה — זה graceful degradation.
    """
    section = doc.sections[0]
    section.different_first_page_header_footer = True

    targets = [
        (section.header, _HEADER_LOGO),
        (section.first_page_header, _HEADER_LOGO),
        (section.footer, _FOOTER_STRIP),
        (section.first_page_footer, _FOOTER_STRIP),
    ]
    for area, image_path in targets:
        # מוחקים פסקאות קיימות כדי לוודא שהתוצאה קבועה
        p = area.paragraphs[0] if area.paragraphs else area.add_paragraph()
        # ניקוי runs קיימים
        for r in list(p.runs):
            r._element.getparent().remove(r._element)
        _apply_rtl_para(p, align_right=False)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if image_path.exists():
            try:
                p.add_run().add_picture(str(image_path), width=Cm(16.5))
            except Exception:
                # asset קיים אבל לא תקין — עדיף ליצור פסקה ריקה מאשר לקרוס
                pass


# ── Paragraph / Run helpers ───────────────────────────────────────────────────

def _apply_rtl_para(p, align_right: bool = True):
    """Put `<w:bidi/>` (and optionally jc=right) on the paragraph."""
    pPr = p._p.get_or_add_pPr()
    # למנוע כפילות אם כבר הוחל
    if pPr.find(qn("w:bidi")) is None:
        bidi = OxmlElement("w:bidi")
        bidi.set(qn("w:val"), "1")
        pPr.append(bidi)
    if align_right:
        jc = pPr.find(qn("w:jc"))
        if jc is None:
            jc = OxmlElement("w:jc")
            pPr.append(jc)
        jc.set(qn("w:val"), "right")


def _apply_rtl_run(run, font_size: int, bold: bool, underline: bool = False):
    """Force RTL + David + size/bold/underline on a run."""
    run.bold = bold
    rPr = run._r.get_or_add_rPr()

    rtl = OxmlElement("w:rtl")
    rPr.append(rtl)

    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), HEBREW_FONT)
    rFonts.set(qn("w:hAnsi"), HEBREW_FONT)
    rFonts.set(qn("w:cs"), HEBREW_FONT)
    rPr.append(rFonts)

    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(font_size * 2))
    rPr.append(sz)
    szCs = OxmlElement("w:szCs")
    szCs.set(qn("w:val"), str(font_size * 2))
    rPr.append(szCs)

    if bold:
        b = OxmlElement("w:b")
        rPr.append(b)
        bCs = OxmlElement("w:bCs")
        rPr.append(bCs)

    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        rPr.append(u)


def add_para(container, text: str, bold: bool = False,
             font_size: int = FONT_BODY,
             space_before: float = 0, space_after: float = 6) -> any:
    """Add an RTL paragraph to a document, cell, or header/footer."""
    p = container.add_paragraph()
    _apply_rtl_para(p)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    if text:
        run = p.add_run(text)
        _apply_rtl_run(run, font_size, bold)
    return p


def add_heading(doc, text: str, level: int = 1) -> any:
    """RTL heading: bold + underline, body-size font, no Word Heading style.

    הסטיילים המובנים של python-docx (Heading 1/2) מוסיפים גופן כחול גדול.
    במקור, כותרות הסעיפים פשוט מודגשות וקו תחתון, באותו גודל גופן של הטקסט.
    הפרמטר `level` נשמר לתאימות אבל אינו משפיע יותר על גודל הגופן.
    """
    p = doc.add_paragraph()
    _apply_rtl_para(p)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(6)
    if text:
        run = p.add_run(text)
        _apply_rtl_run(run, FONT_BODY, bold=True, underline=True)
    return p


def add_bullet(doc, text: str, font_size: int = FONT_BODY) -> any:
    """RTL bullet paragraph: '• ' prefix in a single RTL run.

    הסיבה שאנחנו לא משתמשים ב-numbering אמיתי: python-docx לא חושף
    הגדרת numbering ידידותית, ובמיוחד לא RTL numbering. ה-glyph של ה-bullet
    כתוב ישירות ברן עם דגל RTL — Word יראה אותו בצד ימין כי הפסקה בעצמה
    bidi+jc=right וה-run עם <w:rtl/>.
    """
    p = doc.add_paragraph()
    _apply_rtl_para(p)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.right_indent = Cm(0.5)

    run = p.add_run(f"• {text}")
    _apply_rtl_run(run, font_size, False)
    return p


# ── Tables ────────────────────────────────────────────────────────────────────

def _set_table_bidi(table):
    """Add `<w:bidiVisual/>` to a table so columns flip to RTL order."""
    tblPr = table._element.find(qn("w:tblPr"))
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        table._element.insert(0, tblPr)
    if tblPr.find(qn("w:bidiVisual")) is None:
        bidi = OxmlElement("w:bidiVisual")
        tblPr.append(bidi)


def set_cell_shading(cell, fill_hex: str):
    """Fill a cell with a solid background color (e.g. 'D9D9D9')."""
    tcPr = cell._tc.get_or_add_tcPr()
    # מסירים shading קיים אם יש
    for el in tcPr.findall(qn("w:shd")):
        tcPr.remove(el)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    tcPr.append(shd)


def set_cell(cell, text: str, bold: bool = False, font_size: int = FONT_BODY):
    """Write RTL text into a table cell (clears existing content)."""
    tc = cell._tc
    for p_el in tc.findall(qn("w:p")):
        tc.remove(p_el)

    p_el = OxmlElement("w:p")
    tc.append(p_el)

    pPr = OxmlElement("w:pPr")
    p_el.insert(0, pPr)
    bidi_el = OxmlElement("w:bidi")
    bidi_el.set(qn("w:val"), "1")
    pPr.append(bidi_el)
    jc_el = OxmlElement("w:jc")
    jc_el.set(qn("w:val"), "right")
    pPr.append(jc_el)

    r_el = OxmlElement("w:r")
    p_el.append(r_el)

    rPr = OxmlElement("w:rPr")
    r_el.insert(0, rPr)

    rtl_el = OxmlElement("w:rtl")
    rPr.append(rtl_el)

    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), HEBREW_FONT)
    rFonts.set(qn("w:hAnsi"), HEBREW_FONT)
    rFonts.set(qn("w:cs"), HEBREW_FONT)
    rPr.append(rFonts)

    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(font_size * 2))
    rPr.append(sz)
    szCs = OxmlElement("w:szCs")
    szCs.set(qn("w:val"), str(font_size * 2))
    rPr.append(szCs)

    if bold:
        b_el = OxmlElement("w:b")
        rPr.append(b_el)
        bCs = OxmlElement("w:bCs")
        rPr.append(bCs)

    t_el = OxmlElement("w:t")
    t_el.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t_el.text = text
    r_el.append(t_el)


def make_table(doc, num_rows: int, num_cols: int,
               col_widths_cm: list = None, with_borders: bool = True):
    """Create an RTL table. `with_borders=False` produces a borderless layout."""
    table = doc.add_table(rows=num_rows, cols=num_cols)
    if with_borders:
        table.style = "Table Grid"
    _set_table_bidi(table)
    if col_widths_cm:
        for row in table.rows:
            for i, cell in enumerate(row.cells):
                if i < len(col_widths_cm):
                    cell.width = Cm(col_widths_cm[i])
    return table


# ── Field-line helpers (for borderless detail pages) ──────────────────────────

def add_field_line(doc, label: str, value: str,
                   label_pos_dxa: int = 2636,
                   colon_pos_dxa: int = 3203):
    """Render `<bold-label> [TAB] : [TAB] <value>` as one RTL paragraph.

    משמש בסעיף "פרטי הנכס" כדי להחליף טבלה בפסקאות מסודרות עם tab stops —
    הצורה שבה הדוח המקורי של המשרד בנוי. ה-tab stops קבועים ב-DXA כדי
    שהשורות יושבות זו תחת זו.
    """
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()

    bidi = OxmlElement("w:bidi")
    pPr.append(bidi)
    jc = OxmlElement("w:jc")
    jc.set(qn("w:val"), "right")
    pPr.append(jc)

    tabs = OxmlElement("w:tabs")
    for pos in (label_pos_dxa, colon_pos_dxa):
        t = OxmlElement("w:tab")
        t.set(qn("w:val"), "left")
        t.set(qn("w:pos"), str(pos))
        tabs.append(t)
    pPr.append(tabs)

    p.paragraph_format.space_after = Pt(2)

    r1 = p.add_run(label)
    _apply_rtl_run(r1, FONT_BODY, bold=True)
    r2 = p.add_run("\t:\t")
    _apply_rtl_run(r2, FONT_BODY, bold=False)
    r3 = p.add_run(value)
    _apply_rtl_run(r3, FONT_BODY, bold=False)
    return p
