#!/usr/bin/env python3
"""
Market Research Agent

Researches an industry or niche by reading Reddit, scraping web content,
and gathering reviews/sentiment. Produces a structured analysis of the
top 5 pain points and SaaS opportunities. Saves results to Excel.

Usage:
  python run.py "restaurant industry"
  python run.py "dental practices" --output /path/to/results.xlsx
  python run.py --help
"""

import argparse
import json
import os
import re
import sys
import time
from datetime import datetime
from pathlib import Path

import requests
import trafilatura
import anthropic
import openpyxl
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from duckduckgo_search import DDGS


def log(msg: str):
    print(msg, flush=True)


# ---------------------------------------------------------------------------
# Subreddit discovery
# ---------------------------------------------------------------------------

def get_subreddits(industry: str, client: anthropic.Anthropic) -> dict:
    log(f"→ Identifying relevant subreddits for '{industry}'...")

    msg = client.messages.create(
        model="claude-haiku-4-5",
        max_tokens=500,
        messages=[{
            "role": "user",
            "content": (
                f"List the most relevant Reddit subreddits for researching the '{industry}' industry.\n\n"
                "Two categories:\n"
                "1. operator_subs: subreddits where business owners, operators, professionals discuss their work\n"
                "2. customer_subs: subreddits where customers, clients, or end-users discuss their experiences\n\n"
                'Return ONLY a JSON object, no prose:\n'
                '{"operator_subs": ["sub1", "sub2", ...], "customer_subs": ["sub3", "sub4", ...]}\n\n'
                "Include 4-6 subreddits per category. Real, active subreddits only. No r/ prefix."
            ),
        }],
    )

    text = msg.content[0].text.strip()
    match = re.search(r'\{.*\}', text, re.DOTALL)
    if match:
        try:
            data = json.loads(match.group())
            ops = data.get("operator_subs", [])[:6]
            cust = data.get("customer_subs", [])[:6]
            log(f"  Operator subreddits : {', '.join(ops)}")
            log(f"  Customer subreddits : {', '.join(cust)}")
            return {"operators": ops, "customers": cust}
        except json.JSONDecodeError:
            pass

    log("  Warning: could not parse subreddit list — proceeding without Reddit data")
    return {"operators": [], "customers": []}


# ---------------------------------------------------------------------------
# Reddit scraper (no auth — public JSON API)
# ---------------------------------------------------------------------------

REDDIT_HEADERS = {"User-Agent": "MarketResearchAgent/1.0 (research purposes)"}


def _get_comments(subreddit: str, post_id: str) -> list[str]:
    url = f"https://www.reddit.com/r/{subreddit}/comments/{post_id}.json?limit=6&sort=top"
    try:
        r = requests.get(url, headers=REDDIT_HEADERS, timeout=12)
        if r.status_code != 200:
            return []
        data = r.json()
        comments = []
        for child in data[1]["data"]["children"][:6]:
            body = child.get("data", {}).get("body", "")
            if body and body not in ("[deleted]", "[removed]") and len(body) > 30:
                comments.append(body[:600])
        return comments
    except Exception:
        return []


def scrape_subreddit(subreddit: str, limit: int = 10) -> list[dict]:
    posts = []
    for timeframe in ("month", "year"):
        url = f"https://www.reddit.com/r/{subreddit}/top.json?limit={limit}&t={timeframe}"
        try:
            r = requests.get(url, headers=REDDIT_HEADERS, timeout=15)
            if r.status_code == 429:
                log(f"  Rate-limited on r/{subreddit}, skipping")
                return []
            if r.status_code != 200:
                continue
            children = r.json().get("data", {}).get("children", [])
            for child in children:
                p = child.get("data", {})
                title = p.get("title", "").strip()
                if not title:
                    continue
                text = (p.get("selftext") or "")[:800]
                post_id = p.get("id", "")
                comments = []
                if post_id:
                    comments = _get_comments(subreddit, post_id)
                    time.sleep(0.6)
                posts.append({"title": title, "text": text, "comments": comments})
            if posts:
                break
        except Exception as e:
            log(f"  Warning: r/{subreddit} fetch error — {e}")
    return posts


def gather_reddit(subreddits: dict) -> str:
    log("→ Reading Reddit discussions...")
    all_subs = (
        [(s, "operators/pros") for s in subreddits.get("operators", [])] +
        [(s, "customers/users") for s in subreddits.get("customers", [])]
    )

    sections = []
    for sub, category in all_subs[:10]:
        log(f"  r/{sub} ({category})")
        posts = scrape_subreddit(sub, limit=8)
        if not posts:
            continue
        lines = [f"=== r/{sub} ({category}) ==="]
        for p in posts[:8]:
            entry = f"POST: {p['title']}"
            if p["text"]:
                entry += f"\n{p['text']}"
            if p["comments"]:
                entry += "\nTop comments:\n" + "\n".join(f"  • {c}" for c in p["comments"][:3])
            lines.append(entry)
        sections.append("\n\n".join(lines))
        time.sleep(1.2)

    log(f"  Collected from {len(sections)} subreddits")
    return ("\n\n" + "─" * 60 + "\n\n").join(sections)


# ---------------------------------------------------------------------------
# Web research
# ---------------------------------------------------------------------------

def _scrape_url(url: str) -> str:
    try:
        headers = {"User-Agent": "Mozilla/5.0 (compatible; MarketResearchAgent/1.0)"}
        r = requests.get(url, headers=headers, timeout=14, allow_redirects=True)
        if r.status_code == 200:
            text = trafilatura.extract(
                r.text,
                include_comments=False,
                include_tables=False,
                no_fallback=False,
            )
            return (text or "")[:2500]
    except Exception:
        pass
    return ""


def _ddg_search(query: str, max_results: int = 5) -> list[dict]:
    try:
        with DDGS() as ddgs:
            return [
                {"title": r.get("title", ""), "url": r.get("href", ""), "snippet": r.get("body", "")}
                for r in ddgs.text(query, max_results=max_results)
            ]
    except Exception as e:
        log(f"  Search error for '{query}': {e}")
        return []


def gather_web(industry: str) -> str:
    log(f"→ Searching the web for '{industry}'...")

    queries = [
        f"{industry} biggest problems challenges operators face",
        f"{industry} owner frustrations pain points forum",
        f"{industry} software tools reviews complaints",
        f"{industry} customer complaints bad experience",
        f"{industry} market trends gaps opportunities 2024",
        f"best {industry} software alternatives reddit OR trustpilot",
    ]

    sections = []
    scraped = 0

    for query in queries:
        log(f"  Searching: {query}")
        results = _ddg_search(query, max_results=4)
        for r in results:
            if scraped >= 20:
                break
            url = r.get("url", "")
            title = r.get("title", "")
            snippet = r.get("snippet", "")

            # Always include the snippet
            entry = f"[{title}]\nURL: {url}\n{snippet}"

            # Try full page extraction
            full = _scrape_url(url)
            if full:
                entry += f"\nFull text:\n{full}"

            sections.append(entry)
            scraped += 1
            time.sleep(0.4)

        if scraped >= 20:
            break

    log(f"  Scraped {scraped} web sources")
    return ("\n\n" + "─" * 60 + "\n\n").join(sections)


# ---------------------------------------------------------------------------
# Claude analysis
# ---------------------------------------------------------------------------

ANALYSIS_SCHEMA = """{
  "industry_summary": "2-3 sentence overview",
  "pain_points": [
    {
      "title": "Short descriptive name",
      "description": "Detailed description (2-3 sentences)",
      "acuteness_score": <1-10>,
      "acuteness_explanation": "Why this severity — how urgent / frequent / widespread?",
      "impact_on_operators": "How this hurts business owners, operators, professionals",
      "impact_on_customers": "How this hurts customers, clients, end-users",
      "additional_insights": "Hidden nuances, adjacent trends, what a founder should know",
      "saas_opportunity_score": <1-10>,
      "saas_opportunity_rationale": "Market size, willingness to pay, competitive landscape, feasibility"
    }
  ],
  "overall_saas_verdict": "2-3 sentence verdict on the SaaS opportunity in this industry"
}"""


def analyze(industry: str, reddit: str, web: str, client: anthropic.Anthropic) -> dict:
    log("→ Analyzing with Claude Opus 4.7 (adaptive thinking enabled — may take a moment)...")

    prompt = (
        f"Analyze the '{industry}' industry based on the research data below.\n\n"
        f"## Reddit discussions\n{reddit[:9000]}\n\n"
        f"## Web research\n{web[:9000]}\n\n"
        "─────────────────────────────────────────────────────────\n"
        "Produce a market research report as a single JSON object matching this schema exactly:\n"
        f"{ANALYSIS_SCHEMA}\n\n"
        "Requirements:\n"
        "• Exactly 5 pain points, ordered highest → lowest acuteness_score\n"
        "• Be specific and concrete — ground every claim in the research data\n"
        "• acuteness_score: 1 = minor nuisance, 10 = existential / business-critical\n"
        "• saas_opportunity_score: 1 = crowded/low WTP, 10 = huge unmet need + strong WTP\n"
        "• Output ONLY the JSON — no markdown, no prose before or after"
    )

    result_text = ""
    with client.messages.stream(
        model="claude-opus-4-7",
        max_tokens=4000,
        thinking={"type": "adaptive", "display": "summarized"},
        system=(
            "You are an expert market researcher and SaaS opportunity analyst. "
            "You synthesize qualitative data from forums, reviews, and web content "
            "into actionable, specific insights for founders evaluating market opportunities."
        ),
        messages=[{"role": "user", "content": prompt}],
    ) as stream:
        for text in stream.text_stream:
            result_text += text
            print(".", end="", flush=True)

    print()  # newline

    # Extract JSON (handle optional ```json ... ``` fences)
    code_block = re.search(r'```(?:json)?\s*(.*?)\s*```', result_text, re.DOTALL)
    raw_json = code_block.group(1) if code_block else result_text
    brace_match = re.search(r'\{.*\}', raw_json, re.DOTALL)
    if brace_match:
        try:
            return json.loads(brace_match.group())
        except json.JSONDecodeError as e:
            log(f"  Warning: JSON parse error — {e}")

    log("  Warning: returning raw text")
    return {"raw": result_text}


# ---------------------------------------------------------------------------
# Console report
# ---------------------------------------------------------------------------

def print_report(industry: str, analysis: dict):
    bar = "═" * 70
    thin = "─" * 70

    print(f"\n{bar}")
    print(f"  MARKET RESEARCH: {industry.upper()}")
    print(bar)

    if "raw" in analysis:
        print(analysis["raw"])
        return

    print(f"\n{analysis.get('industry_summary', '')}\n")

    for i, p in enumerate(analysis.get("pain_points", []), 1):
        print(f"\n{thin}")
        print(f"#{i}  {p.get('title', '')}   "
              f"[Acuteness {p.get('acuteness_score', '?')}/10 | "
              f"SaaS opportunity {p.get('saas_opportunity_score', '?')}/10]")
        print(thin)
        print(f"\n{p.get('description', '')}")
        print(f"\nAcuteness: {p.get('acuteness_explanation', '')}")
        print(f"\nOperator impact:\n  {p.get('impact_on_operators', '')}")
        print(f"\nCustomer impact:\n  {p.get('impact_on_customers', '')}")
        print(f"\nAdditional insights:\n  {p.get('additional_insights', '')}")
        print(f"\nSaaS opportunity rationale:\n  {p.get('saas_opportunity_rationale', '')}")

    print(f"\n{bar}")
    print("  OVERALL SaaS VERDICT")
    print(bar)
    print(f"\n{analysis.get('overall_saas_verdict', '')}\n")


# ---------------------------------------------------------------------------
# Excel export
# ---------------------------------------------------------------------------

PAIN_FIELDS = [
    ("title",                    "Title"),
    ("acuteness_score",          "Acuteness (1-10)"),
    ("acuteness_explanation",    "Acuteness Explanation"),
    ("description",              "Description"),
    ("impact_on_operators",      "Impact on Operators"),
    ("impact_on_customers",      "Impact on Customers"),
    ("additional_insights",      "Additional Insights"),
    ("saas_opportunity_score",   "SaaS Score (1-10)"),
    ("saas_opportunity_rationale","SaaS Rationale"),
]

HEADER_COLOR  = "4F46E5"
ALT_ROW_COLOR = "F0F0FF"


def _build_headers() -> list[str]:
    h = ["Industry", "Timestamp", "Industry Summary", "Overall SaaS Verdict"]
    for i in range(1, 6):
        for _, label in PAIN_FIELDS:
            h.append(f"Pain {i} — {label}")
    return h


def save_excel(industry: str, analysis: dict, output_path: Path):
    log(f"→ Saving to {output_path}...")

    if output_path.exists():
        wb = load_workbook(output_path)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Market Research"

        headers = _build_headers()
        ws.append(headers)

        header_font  = Font(bold=True, color="FFFFFF", size=11)
        header_fill  = PatternFill(start_color=HEADER_COLOR, end_color=HEADER_COLOR, fill_type="solid")
        for cell in ws[1]:
            cell.font       = header_font
            cell.fill       = header_fill
            cell.alignment  = Alignment(wrap_text=True, vertical="center")

        ws.row_dimensions[1].height = 30

    # Build data row
    row: list = [
        industry,
        datetime.now().strftime("%Y-%m-%d %H:%M"),
        analysis.get("industry_summary", ""),
        analysis.get("overall_saas_verdict", ""),
    ]

    pain_points = analysis.get("pain_points", [])
    for i in range(5):
        p = pain_points[i] if i < len(pain_points) else {}
        for key, _ in PAIN_FIELDS:
            row.append(p.get(key, ""))

    ws.append(row)

    # Style the new data row
    row_idx = ws.max_row
    fill = PatternFill(start_color=ALT_ROW_COLOR, end_color=ALT_ROW_COLOR, fill_type="solid") \
        if row_idx % 2 == 0 else None
    for cell in ws[row_idx]:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if fill:
            cell.fill = fill

    # Set column widths
    col_widths = [28, 16, 50, 60]  # Industry, Timestamp, Summary, Verdict
    for _ in range(5 * len(PAIN_FIELDS)):
        col_widths.append(45)
    for i, width in enumerate(col_widths, 1):
        col_letter = ws.cell(row=1, column=i).column_letter
        ws.column_dimensions[col_letter].width = width

    ws.freeze_panes = "C2"

    wb.save(output_path)
    log(f"✓ Saved → {output_path}")


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(
        description=(
            "Market research agent. Reads Reddit + web content for an industry, "
            "identifies the top 5 pain points, and evaluates SaaS opportunities."
        )
    )
    parser.add_argument(
        "industry",
        nargs="*",
        help='Industry or niche to research, e.g. "dental practices" or "food trucks"',
    )
    parser.add_argument(
        "--output",
        default=None,
        help="Path to Excel output file (default: market_research.xlsx in agent directory)",
    )
    args = parser.parse_args()

    if not args.industry:
        parser.print_help()
        sys.exit(1)

    api_key = os.environ.get("ANTHROPIC_API_KEY")
    if not api_key:
        print("Error: ANTHROPIC_API_KEY is not set.", file=sys.stderr)
        sys.exit(1)

    industry    = " ".join(args.industry).strip()
    output_path = Path(args.output) if args.output else Path(__file__).parent / "market_research.xlsx"
    client      = anthropic.Anthropic(api_key=api_key)

    print(f"\n{'━'*60}")
    print(f"  Market Research Agent  —  {industry}")
    print(f"{'━'*60}\n")

    subreddits    = get_subreddits(industry, client)
    reddit_data   = gather_reddit(subreddits)
    web_data      = gather_web(industry)
    analysis      = analyze(industry, reddit_data, web_data, client)

    print_report(industry, analysis)
    save_excel(industry, analysis, output_path)

    print(f"\n✓ Done. Results appended to: {output_path}\n")


if __name__ == "__main__":
    main()
